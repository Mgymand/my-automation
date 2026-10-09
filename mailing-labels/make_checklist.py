#!/usr/bin/env python3
"""addresses.json から 送付先確認一覧.xlsx を出力する"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
entries = json.loads((HERE / "addresses.json").read_text(encoding="utf-8"))

wb = Workbook()
ws = wb.active
ws.title = "送付先確認一覧"
headers = ["No", "Excel上の事業所名", "宛名（正式名称）", "運営法人", "郵便番号", "住所",
           "電話番号", "部数", "確認度", "注意点", "出典URL"]
ws.append(headers)
for e in entries:
    ws.append([
        e["no"], e.get("excel_name", ""), e["name"], e.get("corp", ""), e["postal"],
        (e["address1"] + (" " + e["address2"] if e.get("address2") else "")),
        e.get("phone", ""), e.get("copies", ""), e.get("confidence", ""),
        e.get("notes", ""), "\n".join(e.get("sources", [])),
    ])
ws.append([])
ws.append(["合計", f"{len(entries)}事業所", "", "", "", "", "",
           sum(int(e.get("copies") or 0) for e in entries)])

widths = [5, 34, 40, 30, 10, 44, 14, 6, 8, 40, 60]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="DDEBF7")
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    conf = row[8].value
    if conf and conf != "高":
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
ws.freeze_panes = "A2"
wb.save(HERE / "送付先確認一覧.xlsx")
print("送付先確認一覧.xlsx を出力しました")
